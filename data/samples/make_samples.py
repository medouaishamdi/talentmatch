"""Generate fictional demo resumes as PDF and DOCX files.

All people below are invented for the demo. Run from the project root:
    python data/samples/make_samples.py
"""
from pathlib import Path

HERE = Path(__file__).resolve().parent

RESUMES = {
    "sara_haddad_data_scientist": """Sara Haddad
Data Scientist
sara.haddad@example.com | +33 6 12 34 56 78 | linkedin.com/in/sara-haddad-demo | github.com/sarahaddad-demo

Summary
Data scientist with 4 years of experience building machine learning models for retail and fintech.

Experience
Data Scientist, ShopWave (Paris) | Mar 2022 - Present
• Built a demand forecasting model (LightGBM, time series features) that reduced stock-outs by 18%
• Designed and analyzed 30+ A/B tests with product teams; presented results to stakeholders
• Deployed models as REST APIs with FastAPI and Docker on AWS
• Mentored 2 junior analysts on pandas, SQL and statistics

Junior Data Analyst, FinPoint | Sep 2020 - Feb 2022
• Automated weekly reporting with Python and SQL, saving 6 hours per week
• Created Tableau dashboards used by 40 managers

Education
Master of Science in Data Science, Université Paris-Saclay | 2018 - 2020
Bachelor of Science in Mathematics | 2015 - 2018

Skills
Python, pandas, NumPy, scikit-learn, XGBoost, LightGBM, PyTorch, SQL, PostgreSQL, Statistics,
Feature engineering, A/B testing, Time series, NLP, Tableau, Docker, AWS, Git, Spark

Languages
French (native), English (C1), Arabic (B2)
""",
    "yanis_mercier_junior_data": """Yanis Mercier
Junior Data Analyst
yanis.mercier@example.com | +33 7 98 76 54 32

Profile
Recent graduate passionate about data. Hard worker and team player.

Experience
Data Analyst Intern, CityBikes | Jun 2025 - Aug 2025
• Responsible for cleaning data in Excel
• Worked on a dashboard in Power BI
• Helped with SQL queries

Education
Bachelor's degree in Economics and Statistics | 2022 - 2025

Skills
Excel, SQL, Power BI, Python (basics), Statistics, Communication
""",
    "karim_ben_ali_backend": """Karim Ben Ali
Backend Developer (Python)
karim.benali@example.com | +216 22 333 444 | github.com/karimbenali-demo

Experience
Backend Developer, Cloudly | Jan 2023 - Present
• Developed REST APIs with Django and FastAPI serving 1.2M requests/day
• Reduced p95 latency by 35% by adding Redis caching and optimizing PostgreSQL queries
• Containerized 12 services with Docker and deployed them on Kubernetes through GitLab CI
• Wrote unit tests with pytest (coverage from 45% to 85%)

Software Engineering Intern, DataNest | Feb 2022 - Jul 2022
• Built an ETL pipeline in Python and Airflow loading data into PostgreSQL

Education
Engineering degree (Master's level) in Software Engineering | 2019 - 2022

Skills
Python, Django, FastAPI, Flask, PostgreSQL, Redis, Docker, Kubernetes, GitLab CI, Linux, Git,
REST API, Microservices, Celery, Kafka, AWS, pytest
""",
    "lea_dubois_frontend": """Léa Dubois
Frontend Developer
lea.dubois@example.com | +33 6 55 44 33 22 | github.com/leadubois-demo

Experience
Frontend Developer, PixelForge | Sep 2021 - Present
• Built a design system in React and TypeScript used across 6 products
• Improved Lighthouse performance score from 62 to 95 with code splitting (Vite)
• Implemented WCAG accessibility fixes, raising audit score to AA
• Wrote component tests with Jest and React Testing Library

Web Developer, Agence Nova | Jul 2019 - Aug 2021
• Developed 25+ responsive client websites (HTML, CSS, JavaScript, Bootstrap)

Education
Bachelor's degree in Computer Science | 2016 - 2019

Skills
React, TypeScript, JavaScript, Next.js, Redux, HTML, CSS, Tailwind CSS, Responsive design,
Web accessibility, Jest, Vite, Git, Figma, Node.js
""",
    "thomas_laurent_accountant": """Thomas Laurent
Staff Accountant
thomas.laurent@example.com | +33 6 11 22 33 44

Experience
Staff Accountant, Laurent & Associés | May 2020 - Present
• Maintain the general ledger for 15 client companies
• Prepare monthly financial statements and close books within 5 business days
• Process accounts payable and accounts receivable (300+ invoices per month)
• Support external audit and tax preparation

Education
Bachelor's degree in Accounting and Finance | 2016 - 2019

Skills
Accounting, General ledger, Financial reporting, GAAP, IFRS, QuickBooks, Sage, Microsoft Excel,
Payroll, Auditing, Tax preparation
""",
    "ines_moreau_marketing": """Inès Moreau
Digital Marketing Specialist
ines.moreau@example.com | +33 7 12 12 12 12 | linkedin.com/in/ines-moreau-demo

Experience
Digital Marketing Specialist, GreenLeaf | Jan 2022 - Present
• Grew organic traffic by 120% in 12 months through SEO and content marketing
• Managed a €15k/month Google Ads budget with a 4.2x ROAS
• Ran social media marketing campaigns on Instagram and LinkedIn (+8,000 followers)

Marketing Assistant, Bloom Studio | Sep 2020 - Dec 2021
• Wrote newsletters for 20,000 subscribers with Mailchimp

Education
Master's degree in Digital Marketing | 2018 - 2020

Skills
SEO, Google Ads, Social media marketing, Content marketing, Copywriting, Email marketing,
Google Analytics, HubSpot, Adobe Photoshop, Microsoft Excel
""",
}


def write_pdf(path: Path, text: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    lines = text.strip().splitlines()
    story.append(Paragraph(lines[0], styles["Title"]))
    for line in lines[1:]:
        if not line.strip():
            story.append(Spacer(1, 6))
        elif len(line.split()) <= 3 and not line.startswith("•") and "|" not in line and "," not in line:
            story.append(Paragraph(line, styles["Heading3"]))
        else:
            story.append(Paragraph(line.replace("&", "&amp;"), styles["BodyText"]))
    SimpleDocTemplate(str(path), pagesize=A4, title=lines[0]).build(story)


def write_docx(path: Path, text: str) -> None:
    import docx

    document = docx.Document()
    lines = text.strip().splitlines()
    document.add_heading(lines[0], level=0)
    for line in lines[1:]:
        if len(line.split()) <= 3 and line.strip() and not line.startswith("•") and "|" not in line and "," not in line:
            document.add_heading(line, level=2)
        elif line.strip():
            document.add_paragraph(line)
    document.save(path)


if __name__ == "__main__":
    for i, (name, text) in enumerate(RESUMES.items()):
        if i % 2 == 0:
            write_pdf(HERE / f"{name}.pdf", text)
        else:
            write_docx(HERE / f"{name}.docx", text)
    print(f"Wrote {len(RESUMES)} sample resumes to {HERE}")
