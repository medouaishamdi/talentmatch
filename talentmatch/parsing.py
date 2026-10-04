"""Turn resume / job files into clean text and split resumes into sections."""
from __future__ import annotations

import io
import re
from pathlib import Path

import ftfy

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

BULLETS = "•●▪■◦‣∙·➢➤►✓✔❖⦿*"

# Section headings in English and French -> canonical section name
SECTION_PATTERNS: dict[str, list[str]] = {
    "summary": [
        r"summary", r"professional summary", r"profile", r"professional profile", r"objective",
        r"career objective", r"about me", r"executive summary", r"profil", r"r[ée]sum[ée]",
        r"[àa] propos", r"objectif",
    ],
    "experience": [
        r"(?:work |professional |employment |relevant |career |additional |other |industry )?experience",
        r"work history",
        r"employment(?: history)?", r"professional background", r"career history",
        r"exp[ée]riences?(?: professionnelles?)?", r"parcours professionnel", r"stages?",
        r"internships?",
    ],
    "education": [
        r"education(?:al background)?", r"academic(?: background| qualifications)?",
        r"qualifications", r"education and training", r"formations?", r"[ée]tudes", r"dipl[ôo]mes?",
        r"parcours acad[ée]mique", r"education details",
    ],
    "skills": [
        r"(?:technical |core |key |professional |it )?skills", r"skill highlights", r"skills? set",
        r"competencies", r"core competencies", r"technical proficiencies", r"technologies",
        r"tools(?: and technologies)?", r"comp[ée]tences(?: techniques)?", r"savoir-faire",
        r"areas of expertise", r"expertise", r"highlights", r"soft skills", r"personal qualities",
        r"qualit[ée]s",
    ],
    "projects": [r"(?:academic |personal |key )?projects", r"projets?(?: personnels| acad[ée]miques)?"],
    "certifications": [
        r"certifications?", r"certificates?", r"licenses?(?: (?:and|&) certifications)?",
        r"accreditations", r"certificats?",
    ],
    "languages": [r"languages?", r"langues?"],
    "other": [
        r"interests", r"hobbies", r"activities", r"awards", r"honou?rs", r"accomplishments",
        r"achievements", r"volunteer(?:ing| experience)?", r"publications", r"references",
        r"additional information", r"centres? d'int[ée]r[êe]ts?", r"loisirs", r"affiliations",
    ],
}

_HEADING_RES = {
    name: re.compile(r"^\W*(?:" + "|".join(pats) + r")\W*$", re.IGNORECASE)
    for name, pats in SECTION_PATTERNS.items()
}


class UnsupportedFileError(ValueError):
    pass


def extract_text(data: bytes, filename: str) -> str:
    """Extract raw text from PDF, DOCX or plain-text bytes."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif ext in {".txt", ".md"}:
        text = data.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileError(
            f"Unsupported file type '{ext}'. Use one of: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    text = clean_text(text)
    if len(text.split()) < 20:
        raise UnsupportedFileError(
            "Could not read enough text from this file. If it is a scanned PDF (an image), "
            "export it as a text PDF or DOCX."
        )
    return text


def clean_text(text: str) -> str:
    """Fix broken encodings, normalize bullets and whitespace, keep line breaks."""
    text = ftfy.fix_text(text).replace("﻿", "")
    text = re.sub(f"[{re.escape(BULLETS)}]", "\n• ", text)
    text = re.sub(r"_{3,}|-{4,}|={3,}", " ", text)
    text = re.sub(r"[ \t ]+", " ", text)
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(line for line in lines if line and line != "•")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def heading_of(line: str) -> str | None:
    """Return the canonical section name if this line looks like a section heading."""
    stripped = line.strip().rstrip(":").strip()
    if not stripped or len(stripped) > 45 or len(stripped.split()) > 5:
        return None
    for name, rx in _HEADING_RES.items():
        if rx.match(stripped):
            return name
    return None


def split_sections(text: str) -> dict[str, str]:
    """Split resume text into canonical sections. Text before the first heading -> 'header'."""
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in text.splitlines():
        name = heading_of(line)
        if name:
            current = name
            sections.setdefault(current, [])
            continue
        # Inline headings such as "Skills: Python, SQL, ..."
        m = re.match(r"^([A-Za-zÀ-ÿ' ]{3,30}):\s+(.{3,})$", line)
        if m and heading_of(m.group(1)):
            # Inside a skills block, "Languages: Python, C" or "Tools: Git" are skill sub-lists
            if current != "skills":
                current = heading_of(m.group(1))
            sections.setdefault(current, []).append(m.group(2))
            continue
        sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}
